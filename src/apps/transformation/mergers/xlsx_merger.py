import io

import openpyxl

from apps.transformation.interfaces.merger import MergerInterface
from apps.transformation.interfaces.storage import StorageInterface

from .registry import register_merger


@register_merger
class XlsxMerger(MergerInterface):
    format = "xlsx"

    def __init__(self, storage: StorageInterface):
        self._storage = storage

    def merge(self, keys: list[str], metadata: dict) -> bytes:
        # Each chunk file can have multiple named sheets now (per-mapping
        # sheet_name, replacing the old single metadata.sheet_name) -- merge
        # is genuinely three levels deep: chunk file -> sheet within it ->
        # row within that sheet. Kept as one straightforward nested loop
        # rather than split across helpers, since that's the real shape of
        # the work, not incidental structure.
        wb_out = openpyxl.Workbook()
        wb_out.remove(wb_out.active)  # replaced by named sheets below
        sheets_out: dict = {}
        header_written: set = set()

        for key in keys:
            wb_in = openpyxl.load_workbook(
                io.BytesIO(self._storage.get_object(key).read())
            )
            for sheet_name in wb_in.sheetnames:
                # Not setdefault(): its default argument is evaluated eagerly
                # regardless of whether the key already exists, which would
                # call create_sheet() on every chunk -- openpyxl then
                # silently renames the collision ("Identity" -> "Identity1")
                # instead of reusing the first sheet, splitting one logical
                # sheet across chunks into several in the output.
                if sheet_name not in sheets_out:
                    sheets_out[sheet_name] = wb_out.create_sheet(title=sheet_name)
                ws_out = sheets_out[sheet_name]
                for i, row in enumerate(wb_in[sheet_name].iter_rows(values_only=True)):
                    if i == 0:
                        if sheet_name in header_written:
                            continue
                        header_written.add(sheet_name)
                    ws_out.append(list(row))

        buf = io.BytesIO()
        wb_out.save(buf)
        buf.seek(0)
        return buf.read()

import io

import openpyxl

from apps.transformation.mergers.xlsx_merger import XlsxMerger


class FakeStorage:
    def __init__(self, objects: dict):
        self._objects = objects

    def get_object(self, key):
        return io.BytesIO(self._objects[key])


def _xlsx_bytes(sheets: dict):
    """sheets: {sheet_name: [row, row, ...]} -- mirrors what _flush_xlsx
    (ChunkProcessingService) actually writes per chunk: one or more named
    sheets, not a single unnamed active sheet."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_sheet(data: bytes, sheet_name: str):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return list(wb[sheet_name].iter_rows(values_only=True))


class TestXlsxMerger:
    def test_writes_header_once_and_appends_rows_from_every_chunk(self):
        storage = FakeStorage(
            {
                "c1": _xlsx_bytes(
                    {"Export": [["Name", "Area"], ["Acme", 12.5]]}
                ),
                "c2": _xlsx_bytes({"Export": [["Name", "Area"], ["Beta", 8.0]]}),
            }
        )
        merger = XlsxMerger(storage)

        rows = _read_sheet(merger.merge(["c1", "c2"], metadata={}), "Export")

        assert rows == [("Name", "Area"), ("Acme", 12.5), ("Beta", 8.0)]

    def test_sheet_names_come_from_the_chunk_files_themselves(self):
        # Sheet titles are no longer sourced from metadata.sheet_name -- each
        # chunk file already carries its own real sheet name(s), written by
        # _flush_xlsx from the per-mapping sheet_name.
        storage = FakeStorage({"c1": _xlsx_bytes({"Farms": [["Name"], ["Acme"]]})})
        merger = XlsxMerger(storage)

        result = merger.merge(["c1"], metadata={})

        wb = openpyxl.load_workbook(io.BytesIO(result))
        assert wb.sheetnames == ["Farms"]

    def test_multiple_sheets_are_merged_independently_and_matched_by_name(self):
        storage = FakeStorage(
            {
                "c1": _xlsx_bytes(
                    {
                        "Identity": [["Name"], ["Acme"]],
                        "Metrics": [["Area"], [12.5]],
                    }
                ),
                "c2": _xlsx_bytes(
                    {
                        "Identity": [["Name"], ["Beta"]],
                        "Metrics": [["Area"], [8.0]],
                    }
                ),
            }
        )
        merger = XlsxMerger(storage)

        result = merger.merge(["c1", "c2"], metadata={})

        wb = openpyxl.load_workbook(io.BytesIO(result))
        assert wb.sheetnames == ["Identity", "Metrics"]
        assert _read_sheet(result, "Identity") == [("Name",), ("Acme",), ("Beta",)]
        assert _read_sheet(result, "Metrics") == [("Area",), (12.5,), (8.0,)]
